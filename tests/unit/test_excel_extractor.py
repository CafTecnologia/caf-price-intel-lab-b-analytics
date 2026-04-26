from pathlib import Path

from openpyxl import Workbook

from procurement_core.extractors.excel import ExcelExtractor
from procurement_core.models.enums import FileType
from procurement_core.schemas.extraction import DocumentProfile


def test_excel_extractor_reads_tabular_items(tmp_path: Path) -> None:
    file_path = tmp_path / "anexo_economico.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Descripción", "Cantidad", "Unidad", "Valor Unitario", "Valor Total", "Lote"])
    sheet.append(["Guantes de nitrilo talla M caja x 100", 10, "caja", 32000, 320000, "Lote 1"])
    workbook.save(file_path)

    batch = ExcelExtractor().extract(file_path)

    assert batch.document_kind.value in {"economic_annex", "price_list", "unknown"}
    assert len(batch.items) == 1
    assert batch.items[0].raw_description == "Guantes de nitrilo talla M caja x 100"


def test_excel_extractor_handles_multilevel_headers_and_reference_prices(tmp_path: Path) -> None:
    file_path = tmp_path / "oferta_economica.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Formato Oferta Económica"
    sheet.append([None] * 15)
    sheet.append([None, None, None, None, None, None, "DISERMED", "FEDEPFAGROS", "AGROTECH", "VALOR PROMEDIO - UNITARIO", "VALOR PROMEDIO TOTAL", "MENOR VALOR - UNITARIO", "MENOR VALOR -  TOTAL", "PRECIO TECHO UNITARIO - CON IVA", "PRECIO TECHO TOTAL - CON IVA"])
    sheet.append([None, "ITEM", "ELEMENTO", "UNIDAD", "DESCRIPCIÓN", "CANTIDAD", "VALOR UNITARIO", "VALOR UNITARIO", "VALOR UNITARIO", None, None, None, None, None, None])
    sheet.append([None, 1, "Resma", "Resma", "Resma de papel alcalino 500 hojas", 24, 22275, 26325, 20250, 22950, 550800, 20250, 486000, 21600, 518400])
    workbook.save(file_path)

    batch = ExcelExtractor().extract(file_path)

    assert len(batch.items) == 1
    assert batch.items[0].raw_description == "Resma de papel alcalino 500 hojas"
    assert batch.items[0].raw_unit == "Resma"
    assert batch.items[0].raw_quantity == "24"
    assert batch.items[0].raw_unit_price == "21600"
    assert batch.items[0].raw_total == "518400"


def test_excel_extractor_prefers_average_columns_and_filters_irrelevant_sheets(tmp_path: Path) -> None:
    file_path = tmp_path / "estudio_mercado.xlsx"
    workbook = Workbook()

    market_sheet = workbook.active
    market_sheet.title = "ESTUDIO DE MERCADO"
    market_sheet.append([None] * 18)
    market_sheet.append([None] * 18)
    market_sheet.append([None] * 18)
    market_sheet.append([None] * 18)
    market_sheet.append([None] * 18)
    market_sheet.append([None] * 18)
    market_sheet.append([None] * 18)
    market_sheet.append([None, None, None, None, None, None, "ELECTRIPLAZA", None, None, None, "REDESUR", None, None, None, "PROMEDIO SENA", None, None, None])
    market_sheet.append(["ITEM", "CODIGO UNSPSC", "NOMBRE DEL ELEMENTO", "Unidad de medida", "Cantidad", "ESPECIFICACION", "VALOR UNITARIO SIN IVA", "IVA", "VALOR IVA", "VALOR TOTAL", "VALOR UNITARIO SIN IVA", "IVA", "VALOR IVA", "VALOR TOTAL", "VALOR UNITARIO SIN IVA", "IVA", "VALOR IVA", "VALOR TOTAL"])
    market_sheet.append([1, "111", "ADAPTADOR TERMINAL", "UNIDAD", 10, "Compra", 640, 0.19, 121.6, 7616, 720, 0.19, 136.8, 8568, 680, 0.19, 129.2, 8092])
    market_sheet.append([2, "222", "BREAKER 20A", "UNIDAD", 20, "Compra", 16000, 0.19, 3040, 380800, 17320, 0.19, 3290.8, 412216, 16660, 0.19, 3165.4, 396508])
    market_sheet.append([3, "333", "CABLE DUPLEX", "UNIDAD", 5, "Compra", 120000, 0.19, 22800, 714000, 130000, 0.19, 24700, 773500, 125000, 0.19, 23750, 743750])

    support_sheet = workbook.create_sheet("Hoja3")
    support_sheet.append([None] * 12)
    support_sheet.append([None] * 12)
    support_sheet.append([None] * 12)
    support_sheet.append([None] * 12)
    support_sheet.append([None, None, None, None, None, None, "ITEM", "CODIGO UNSPSC", "DESCRIPCION", "UNIDAD DE MEDIDA", "CANTIDAD", None])
    support_sheet.append([None, None, None, None, None, None, 1, "111", "ADAPTADOR TERMINAL", "UND", 12, None])
    support_sheet.append([None, None, None, None, None, None, 2, "222", "BREAKER 20A", "UND", 25, None])
    support_sheet.append([None, None, None, None, None, None, 3, "333", "CABLE DUPLEX", "UND", 7, None])

    noise_sheet = workbook.create_sheet("Hoja1")
    noise_sheet.append(["DESCRIPCION", "CANTIDAD", "UNIDAD", "VALOR UNITARIO", "VALOR TOTAL"])
    noise_sheet.append(["Reactivo de microbiologia", 1, "Unidad", 999999, 999999])

    workbook.save(file_path)

    batch = ExcelExtractor().extract(file_path)

    assert len(batch.items) == 6
    assert {item.sheet_name for item in batch.items} == {"ESTUDIO DE MERCADO", "Hoja3"}
    market_items = [item for item in batch.items if item.sheet_name == "ESTUDIO DE MERCADO"]
    support_items = [item for item in batch.items if item.sheet_name == "Hoja3"]
    assert market_items[0].raw_item_number == "1"
    assert market_items[0].raw_unit_price == "680"
    assert market_items[0].raw_total == "8092"
    assert market_items[0].evidence["row_role"] == "official_item"
    assert market_items[0].evidence["sheet_role"] == "official"
    assert market_items[1].raw_unit_price == "16660"
    assert market_items[1].raw_total == "396508"
    assert support_items[0].evidence["row_role"] == "support_item"
    assert support_items[0].evidence["sheet_role"] == "support"


def test_excel_extractor_uses_original_profile_name_for_classification(tmp_path: Path) -> None:
    file_path = tmp_path / "6c6b1ade8c33aecf.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ESTUDIO DE MERCADO"
    sheet.append(["ITEM", "NOMBRE DEL ELEMENTO", "Unidad de medida", "Cantidad", "VALOR UNITARIO SIN IVA", "VALOR TOTAL"])
    sheet.append([1, "ADAPTADOR TERMINAL", "UNIDAD", 10, 680, 8092])
    workbook.save(file_path)

    profile = DocumentProfile(
        file_name="Estudio de Mercado Electricidad.xlsx",
        file_type=FileType.EXCEL,
        route="excel",
    )

    batch = ExcelExtractor().extract(file_path, profile=profile)

    assert batch.document_kind.value == "market_study"


def test_excel_extractor_preserves_tax_rate_when_excel_formats_it_as_percent(tmp_path: Path) -> None:
    file_path = tmp_path / "tributos.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ESTUDIO DE MERCADO"
    sheet.append([None] * 18)
    sheet.append([None, None, None, None, None, None, "ELECTRIPLAZA", None, None, None, "REDESUR", None, None, None, "PROMEDIO SENA", None, None, None])
    sheet.append(["ITEM", "CODIGO UNSPSC", "NOMBRE DEL ELEMENTO", "Unidad de medida", "Cantidad", "ESPECIFICACION", "VALOR UNITARIO SIN IVA", "IVA", "VALOR IVA", "VALOR TOTAL", "VALOR UNITARIO SIN IVA", "IVA", "VALOR IVA", "VALOR TOTAL", "VALOR UNITARIO SIN IVA", "IVA", "VALOR IVA", "VALOR TOTAL"])
    sheet.append([1, "111", "ADAPTADOR TERMINAL", "UNIDAD", 10, "Compra", 640, 0.19, 121.6, 7616, 720, 0.19, 136.8, 8568, 680, 0.19, 129.2, 8092])
    sheet["H4"].number_format = "0%"
    sheet["L4"].number_format = "0%"
    sheet["P4"].number_format = "0%"
    workbook.save(file_path)

    batch = ExcelExtractor().extract(file_path)

    assert batch.items[0].raw_tax_note == "0.19"


def test_excel_extractor_accepts_reference_price_column_and_carries_unit_down(tmp_path: Path) -> None:
    file_path = tmp_path / "acacia.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Table 1"
    sheet.append(["Item", "Cant.", "Und/Med", "Descripción", "Precio ref Promedio Aritm.", None])
    sheet.append([1, 12, "Unidad", "CANDADO MEDIANO", 840000, None])
    sheet.append([2, 10, None, "CANDADO MEDIANO CON CLAVE", 500000, None])
    sheet.append([3, 150, None, "CARPETA SEGURIDAD OFICIO", 1200000, None])
    workbook.save(file_path)

    batch = ExcelExtractor().extract(file_path)

    assert len(batch.items) == 3
    assert [item.raw_item_number for item in batch.items] == ["1", "2", "3"]
    assert [item.raw_total for item in batch.items] == ["840000", "500000", "1200000"]
    assert [item.raw_unit_price for item in batch.items] == [None, None, None]
    assert [item.raw_unit for item in batch.items] == ["Unidad", "Unidad", "Unidad"]
    assert all(item.evidence["row_role"] == "official_item" for item in batch.items)
    assert all(item.evidence["financial_basis_role"] == "total" for item in batch.items)
    assert batch.items[1].evidence["carried_fields"] == {"unit": "Unidad"}


def test_excel_extractor_prefers_richer_second_header_row_over_partial_first_row(tmp_path: Path) -> None:
    file_path = tmp_path / "candelaria.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "EST MERCADO"
    sheet.append(["No", "Descripción", "Cantidad", "PROMEDIO (IVA Incluido)"])
    sheet.append([None, None, None, "Vr Unitario"])
    sheet.append([1, "TONER LEXMAR MX622", 1, 2451307.5766666667])
    sheet.append([2, "TONER HP CE278A (78A)", 1, 624424.3366666668])
    workbook.save(file_path)

    batch = ExcelExtractor().extract(file_path)

    assert len(batch.items) == 2
    assert [item.raw_item_number for item in batch.items] == ["1", "2"]
    assert [item.raw_unit_price for item in batch.items] == ["2451307.576667", "624424.336667"]
    assert [item.raw_total for item in batch.items] == [None, None]
