from __future__ import annotations

from pathlib import Path

from docx import Document

from procurement_core.extractors.base import BaseExtractor
from procurement_core.extractors.helpers import map_headers, row_to_candidate
from procurement_core.models.enums import ExtractionMethod, FileType
from procurement_core.schemas.extraction import ExtractionBatch, PageExtraction
from procurement_core.services.classification import classify_document_from_text


class WordExtractor(BaseExtractor):
    file_type = FileType.WORD
    route_name = "word"

    def extract(self, file_path: Path, profile=None) -> ExtractionBatch:
        document = Document(str(file_path))
        text_fragments = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        pages = [PageExtraction(page_number=1, extracted_text="\n".join(text_fragments), metadata={"source": "python-docx"})]
        items = []

        for table_index, table in enumerate(document.tables):
            matrix = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if len(matrix) < 2:
                continue
            header_map = map_headers(matrix[0])
            if "description" not in header_map:
                continue
            for row_index, row in enumerate(matrix[1:], start=2):
                candidate = row_to_candidate(
                    row,
                    header_map,
                    page_number=1,
                    row_index=row_index,
                    method=ExtractionMethod.WORD_TABLE,
                )
                if candidate:
                    candidate.section = f"table_{table_index + 1}"
                    items.append(candidate)

        file_name_for_classification = profile.file_name if profile is not None else file_path.name
        document_kind = classify_document_from_text(file_name_for_classification, "\n".join(text_fragments))
        return ExtractionBatch(document_kind=document_kind, pages=pages, items=items, metadata={"table_count": len(document.tables)})
