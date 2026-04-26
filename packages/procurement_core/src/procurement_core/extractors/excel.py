from __future__ import annotations

from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

from procurement_core.extractors.base import BaseExtractor
from procurement_core.extractors.helpers import looks_numericish, row_to_candidate, slug_header, value_from_row
from procurement_core.models.enums import ExtractionMethod, FileType
from procurement_core.schemas.extraction import ExtractionBatch, PageExtraction
from procurement_core.services.item_identity import canonicalize_description
from procurement_core.services.classification import classify_document_from_text


def _excel_column_name(index: int) -> str:
    label = ""
    current = index
    while current > 0:
        current, remainder = divmod(current - 1, 26)
        label = chr(65 + remainder) + label
    return label


def _decimal_places_from_format(number_format: str | None) -> int | None:
    if not number_format:
        return None
    normalized = str(number_format).split(";")[0]
    if "." in normalized:
        return len(normalized.rsplit(".", 1)[1].replace("#", "").replace("0", "0"))
    return 0 if "0" in normalized or "#" in normalized else None


def _format_numeric_excel_value(value: int | float, number_format: str | None) -> str:
    decimals = _decimal_places_from_format(number_format)
    normalized_format = str(number_format or "")
    decimal_value = Decimal(str(value))
    if "%" in normalized_format:
        if decimals is None or decimals <= 0:
            decimals = 4
        quantizer = Decimal("1").scaleb(-decimals)
        quantized = decimal_value.quantize(quantizer, rounding=ROUND_HALF_UP)
        return format(quantized.normalize(), "f")
    if decimals is None:
        decimals = 0 if decimal_value == decimal_value.to_integral() else 6
    if decimals <= 0:
        return str(int(decimal_value.quantize(Decimal("1"), rounding=ROUND_HALF_UP)))
    quantizer = Decimal("1").scaleb(-decimals)
    quantized = decimal_value.quantize(quantizer, rounding=ROUND_HALF_UP)
    return format(quantized.normalize(), "f")


def _format_excel_cell(cell) -> object | None:
    value = cell.value
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return _format_numeric_excel_value(value, getattr(cell, "number_format", None))
    return value


def _is_probable_total_label(value: str) -> bool:
    normalized = slug_header(value)
    return normalized in {"valor total", "subtotal", "total", "total general"} or normalized.startswith("subtotal ")


def _build_composite_headers(values: list[list[object]], header_row_index: int) -> list[str]:
    start_row = max(0, header_row_index - 1)
    composite: list[str] = []
    width = max(len(row) for row in values[: header_row_index + 1]) if values[: header_row_index + 1] else 0
    normalized_rows: list[list[object]] = []
    for row_index in range(start_row, header_row_index + 1):
        source_row = list(values[row_index])
        filled_row: list[object] = []
        last_text: object | None = None
        for col in range(width):
            cell = source_row[col] if col < len(source_row) else None
            if cell not in (None, ""):
                last_text = cell
                filled_row.append(cell)
            else:
                filled_row.append(last_text)
        normalized_rows.append(filled_row)
    for col in range(width):
        parts: list[str] = []
        for row in normalized_rows:
            cell = row[col] if col < len(row) else None
            text = slug_header(cell)
            if text and text not in parts:
                parts.append(text)
        composite.append(" ".join(parts).strip())
    return composite


def _match_best(
    headers: list[str],
    candidates: list[str],
    *,
    prefer_terms: list[str] | None = None,
    avoid_terms: list[str] | None = None,
) -> int | None:
    best_index: int | None = None
    best_score = -1
    prefer_terms = prefer_terms or []
    avoid_terms = avoid_terms or []
    for index, header in enumerate(headers):
        if not header:
            continue
        score = 0
        for candidate_index, candidate in enumerate(candidates):
            priority_bonus = max(0, (len(candidates) - candidate_index) * 2)
            if header == candidate:
                score = max(score, 100 + len(candidate) + priority_bonus)
            elif candidate in header:
                score = max(score, 60 + len(candidate) + priority_bonus)
        if score <= 0:
            continue
        for term in prefer_terms:
            if term in header:
                score += 35
        for term in avoid_terms:
            if term in header:
                score -= 20
        if score > best_score:
            best_score = score
            best_index = index
    return best_index


def _match_item_number(headers: list[str]) -> int | None:
    blocked_terms = {"total", "unitario", "iva", "precio", "valor"}
    exact_aliases = {
        "no",
        "n",
        "nro",
        "numero",
        "item",
        "no item",
        "numero item",
        "item no",
        "item nro",
        "# item",
        "item #",
    }
    for index, header in enumerate(headers):
        if not header:
            continue
        parts = set(header.split())
        if parts & blocked_terms:
            continue
        if header in exact_aliases:
            return index
        if header.startswith("no ") or header.startswith("numero "):
            return index
        if " item " in f" {header} ":
            return index
    return None


def _header_map_from_headers(headers: list[str]) -> dict[str, int] | None:
    description_idx = _match_best(
        headers,
        [
            "descripcion tecnica",
            "descripcion",
            "nombre del elemento",
            "nombre elemento",
            "nombre",
            "especificaciones",
            "especificacion",
        ],
        prefer_terms=["nombre del elemento", "descripcion"],
    )
    quantity_idx = _match_best(headers, ["cantidad ajustada", "cantidad inicial", "cantidad", "cant"])
    unit_idx = _match_best(headers, ["unidad de medida", "unidad medida", "unidad", "und", "u m"])
    unit_price_idx = _match_best(
        headers,
        [
            "promedio sena valor unitario con iva",
            "promedio sena valor unitario sin iva",
            "promedio sena vr unitario con iva",
            "promedio sena vr unitario",
            "precio techo unitario",
            "valor unitario con iva",
            "valor unitario sin iva",
            "valor unitario",
            "vr unitario con iva",
            "vr unitario",
            "v unit",
            "unitario con iva",
            "valor promedio unitario",
            "menor valor unitario",
        ],
        prefer_terms=["promedio", "sena", "referencia", "con iva"],
        avoid_terms=["electriplaza", "redesur"],
    )
    reference_amount_idx = _match_best(
        headers,
        [
            "precio ref promedio aritm",
            "precio ref promedio aritmetico",
            "precio referencia promedio",
            "precio referencia",
            "precio ref",
        ],
        prefer_terms=["precio", "referencia", "promedio", "aritm"],
    )
    total_idx = _match_best(
        headers,
        [
            "promedio sena valor total",
            "promedio sena total item con iva",
            "total item con iva",
            "precio techo total",
            "valor techo total",
            "valor promedio total",
            "menor valor total",
            "vr total",
            "valor total",
            "subtotal",
            "total",
        ],
        prefer_terms=["promedio", "sena", "referencia", "con iva"],
        avoid_terms=["electriplaza", "redesur"],
    )
    tax_idx = _match_best(
        headers,
        [
            "promedio sena iva",
            "iva",
            "valor iva",
        ],
        prefer_terms=["promedio", "sena"],
        avoid_terms=["electriplaza", "redesur"],
    )
    if tax_idx in {unit_price_idx, total_idx}:
        tax_idx = None
    if description_idx is None or quantity_idx is None:
        return None
    header_map = {
        "description": description_idx,
        "quantity": quantity_idx,
    }
    item_number_idx = _match_item_number(headers)
    if item_number_idx is not None:
        header_map["item_number"] = item_number_idx
    if unit_idx is not None:
        header_map["unit"] = unit_idx
    if unit_price_idx is not None:
        header_map["unit_price"] = unit_price_idx
    if total_idx is not None:
        header_map["total"] = total_idx
    if reference_amount_idx is not None and "unit_price" not in header_map and "total" not in header_map:
        header_map["total"] = reference_amount_idx
        header_map["financial_basis_role"] = "total"
        header_map["financial_basis_reason"] = "single_reference_amount_column"
    if tax_idx is not None:
        header_map["tax"] = tax_idx
    lot_idx = _match_best(headers, ["lote", "grupo"])
    if lot_idx is not None:
        header_map["lot"] = lot_idx
    return header_map


def _apply_row_carry_forward(
    row: list[object],
    header_map: dict[str, int],
    carry_state: dict[str, object | None],
) -> tuple[list[object], dict[str, object]]:
    hydrated = list(row)
    carried_fields: dict[str, object] = {}
    for field in ["unit", "lot"]:
        index = header_map.get(field)
        if index is None or index >= len(hydrated):
            continue
        current = hydrated[index]
        if current not in (None, ""):
            carry_state[field] = current
            continue
        carried = carry_state.get(field)
        if carried not in (None, ""):
            hydrated[index] = carried
            carried_fields[field] = carried
    return hydrated, carried_fields


def _detect_header_map(values: list[list[object]]) -> tuple[dict[str, int], int] | tuple[None, None]:
    best_match: tuple[dict[str, int], int, int] | None = None
    for row_index, _row in enumerate(values[:20]):
        direct_headers = [slug_header(value) for value in values[row_index]]
        composite_headers = _build_composite_headers(values, row_index)
        direct_map = _header_map_from_headers(direct_headers)
        composite_map = _header_map_from_headers(composite_headers)
        candidates: list[dict[str, int]] = []
        if direct_map is not None:
            candidates.append(direct_map)
        if composite_map is not None:
            candidates.append(composite_map)
        if direct_map is not None and composite_map is not None:
            merged_map = dict(direct_map)
            for field in ["item_number", "unit_price", "total", "tax", "lot"]:
                if field in composite_map:
                    merged_map[field] = composite_map[field]
            candidates.append(merged_map)

        for candidate in candidates:
            score = len(candidate) * 10
            if "unit_price" in candidate:
                score += 7
            if "total" in candidate:
                score += 7
            if "tax" in candidate:
                score += 2
            if "item_number" in candidate:
                score += 2
            if direct_map is not None and composite_map is not None and candidate is not direct_map:
                score += 3
            if best_match is None or score > best_match[2]:
                best_match = (candidate, row_index, score)
    if best_match is None:
        return None, None
    return best_match[0], best_match[1]


def _infer_flexible_header_map(values: list[list[object]]) -> tuple[dict[str, int], int] | tuple[None, None]:
    for row_index in range(min(len(values), 25)):
        row = values[row_index]
        if not row:
            continue
        non_empty = [(index, value) for index, value in enumerate(row) if value not in (None, "")]
        if len(non_empty) < 3:
            continue
        text_columns = [index for index, value in non_empty if not str(value).strip().replace(".", "").replace(",", "").isdigit()]
        numeric_columns = [index for index, value in non_empty if str(value).strip().replace(".", "").replace(",", "").isdigit()]
        if not text_columns or len(numeric_columns) < 2:
            continue
        description_idx = max(text_columns, key=lambda index: len(str(row[index]).strip()))
        quantity_idx = min(numeric_columns)
        total_idx = max(numeric_columns)
        unit_price_idx = next((index for index in numeric_columns if index not in {quantity_idx, total_idx}), None)
        header_map: dict[str, int] = {
            "description": description_idx,
            "quantity": quantity_idx,
            "total": total_idx,
        }
        if unit_price_idx is not None:
            header_map["unit_price"] = unit_price_idx
        item_number_idx = next(
            (index for index, value in non_empty if slug_header(value) in {"item", "no", "n", "numero"}),
            None,
        )
        if item_number_idx is not None:
            header_map["item_number"] = item_number_idx
        unit_idx = next(
            (
                index
                for index, value in non_empty
                if slug_header(value) in {"unidad", "und", "un", "caja", "kg", "metro", "ml"}
            ),
            None,
        )
        if unit_idx is not None:
            header_map["unit"] = unit_idx
        return header_map, row_index - 1
    return None, None


def _sheet_item_rows(items: list) -> dict[str, list]:
    grouped: dict[str, list] = {}
    for item in items:
        grouped.setdefault(item.sheet_name or "", []).append(item)
    return grouped


def _description_set(sheet_items: list) -> set[str]:
    result: set[str] = set()
    for item in sheet_items:
        normalized = canonicalize_description(item.raw_description)
        if normalized:
            result.add(normalized)
    return result


def _sheet_priority_score(sheet_name: str, sheet_items: list) -> int:
    lowered = sheet_name.lower()
    priced_rows = len([item for item in sheet_items if item.raw_unit_price and item.raw_total])
    numbered_rows = len([item for item in sheet_items if item.raw_item_number])
    score = len(sheet_items) + priced_rows * 3 + numbered_rows * 2
    if "estudio" in lowered or "mercado" in lowered:
        score += 40
    if "presupuesto" in lowered or "oferta" in lowered:
        score += 24
    return score


def _classify_row_role(
    row: list[object],
    header_map: dict[str, int],
    *,
    has_previous_candidate: bool,
) -> str:
    description = value_from_row(row, header_map, "description")
    item_number = value_from_row(row, header_map, "item_number")
    quantity = value_from_row(row, header_map, "quantity")
    unit = value_from_row(row, header_map, "unit")
    unit_price = value_from_row(row, header_map, "unit_price")
    total = value_from_row(row, header_map, "total")

    has_description = bool(str(description).strip()) if description is not None else False
    has_item_number = bool(str(item_number).strip()) if item_number is not None else False
    has_quantity = looks_numericish(quantity)
    has_unit = bool(str(unit).strip()) if unit is not None else False
    has_unit_price = looks_numericish(unit_price)
    has_total = looks_numericish(total)

    if has_description and _is_probable_total_label(str(description)):
        return "subtotal"
    if (has_item_number or has_quantity) and (has_unit_price or has_total):
        return "official_item"
    if has_item_number and has_quantity and has_description:
        return "support_item"
    if has_item_number and has_quantity and has_unit:
        return "support_item"
    if has_quantity and has_unit and not (has_unit_price or has_total):
        return "support_item"
    if has_description and has_previous_candidate and not any([has_item_number, has_quantity, has_unit_price, has_total]):
        return "description_continuation"
    if has_description:
        return "note"
    return "noise"


def _annotate_sheet_roles(items: list) -> tuple[list, str | None]:
    by_sheet = _sheet_item_rows(items)
    if not by_sheet:
        return items, None

    def _official_score(sheet_name: str, sheet_items: list) -> int:
        lowered = sheet_name.lower()
        official_rows = len([item for item in sheet_items if item.evidence.get("row_role") == "official_item"])
        support_rows = len([item for item in sheet_items if item.evidence.get("row_role") == "support_item"])
        numbered_rows = len([item for item in sheet_items if item.raw_item_number])
        complete_rows = len([item for item in sheet_items if item.raw_unit_price and item.raw_total])
        score = official_rows * 12 + complete_rows * 8 + numbered_rows * 3 + support_rows
        if "estudio" in lowered or "mercado" in lowered:
            score += 35
        if "presupuesto" in lowered or "oferta" in lowered:
            score += 20
        return score

    ranked = sorted(by_sheet, key=lambda sheet: _official_score(sheet, by_sheet[sheet]), reverse=True)
    winner = ranked[0]
    winner_score = _official_score(winner, by_sheet[winner])
    winner_official_rows = len([item for item in by_sheet[winner] if item.evidence.get("row_role") == "official_item"])
    second_score = _official_score(ranked[1], by_sheet[ranked[1]]) if len(ranked) > 1 else 0
    official_sheet = winner if winner_official_rows > 0 and winner_score >= max(20, second_score + 8) else None

    for item in items:
        sheet_role = "official"
        if official_sheet is not None and (item.sheet_name or "") != official_sheet:
            sheet_role = "support"
        item.evidence["sheet_role"] = sheet_role
    return items, official_sheet


def _extract_sheet_candidates(
    values: list[list[object]],
    *,
    sheet_name: str,
    header_map: dict[str, int],
    header_row_index: int,
    extraction_mode: str,
) -> tuple[list, dict[str, int]]:
    sheet_items = []
    row_stats = {
        "official_item": 0,
        "support_item": 0,
        "description_continuation": 0,
        "note": 0,
        "subtotal": 0,
        "noise": 0,
    }
    previous_candidate = None
    carry_state: dict[str, object | None] = {}

    for row_index, row in enumerate(values[header_row_index + 1 :], start=header_row_index + 2):
        row_list, carried_fields = _apply_row_carry_forward(list(row), header_map, carry_state)
        row_role = _classify_row_role(row_list, header_map, has_previous_candidate=previous_candidate is not None)
        row_stats[row_role] = row_stats.get(row_role, 0) + 1

        if row_role == "description_continuation" and previous_candidate is not None:
            continuation_text = str(value_from_row(row_list, header_map, "description") or "").strip()
            if continuation_text:
                previous_candidate.raw_description = f"{previous_candidate.raw_description} {continuation_text}".strip()
                continuations = list(previous_candidate.evidence.get("continuation_rows") or [])
                continuations.append(
                    {
                        "sheet_name": sheet_name,
                        "row_index": row_index,
                        "cell_range": f"A{row_index}:{_excel_column_name(max(1, len(row_list)))}{row_index}",
                        "text": continuation_text,
                    }
                )
                previous_candidate.evidence["continuation_rows"] = continuations
            continue

        if row_role not in {"official_item", "support_item"}:
            previous_candidate = None if row_role in {"subtotal", "noise"} else previous_candidate
            continue

        candidate = row_to_candidate(
            row_list,
            header_map,
            sheet_name=sheet_name,
            row_index=row_index,
            method=ExtractionMethod.EXCEL_TABLE,
        )
        if candidate is None:
            continue
        candidate.cell_range = f"A{row_index}:{_excel_column_name(max(1, len(row_list)))}{row_index}"
        candidate.evidence["sheet_header_map"] = header_map
        candidate.evidence["excel_extraction_mode"] = extraction_mode
        candidate.evidence["row_role"] = row_role
        if "financial_basis_role" in header_map:
            candidate.evidence["financial_basis_role"] = header_map["financial_basis_role"]
            candidate.evidence["financial_basis_reason"] = header_map.get("financial_basis_reason")
        if carried_fields:
            candidate.evidence["carried_fields"] = carried_fields
        sheet_items.append(candidate)
        previous_candidate = candidate

    return sheet_items, row_stats


def _filter_to_relevant_sheet_cluster(items: list) -> list:
    by_sheet = _sheet_item_rows(items)
    if len(by_sheet) <= 1:
        return items

    description_sets = {sheet: _description_set(sheet_items) for sheet, sheet_items in by_sheet.items()}
    overlap_support: dict[str, int] = {sheet: 0 for sheet in by_sheet}
    for left_sheet, left_descriptions in description_sets.items():
        for right_sheet, right_descriptions in description_sets.items():
            if left_sheet >= right_sheet:
                continue
            overlap = len(left_descriptions & right_descriptions)
            overlap_support[left_sheet] += overlap
            overlap_support[right_sheet] += overlap

    ranked_sheets = sorted(
        by_sheet,
        key=lambda sheet: (_sheet_priority_score(sheet, by_sheet[sheet]) + overlap_support.get(sheet, 0) * 4),
        reverse=True,
    )
    winner = ranked_sheets[0]
    if overlap_support.get(winner, 0) < 3:
        return items

    allowed = {
        sheet
        for sheet in by_sheet
        if sheet == winner or len(description_sets.get(sheet, set()) & description_sets.get(winner, set())) >= 3
    }
    return [item for item in items if (item.sheet_name or "") in allowed]


class ExcelExtractor(BaseExtractor):
    file_type = FileType.EXCEL
    route_name = "excel"

    def extract(self, file_path: Path, profile=None) -> ExtractionBatch:
        workbook = load_workbook(filename=file_path, data_only=True)
        pages: list[PageExtraction] = []
        items = []
        text_fragments: list[str] = []
        sheet_stats: dict[str, dict[str, int]] = {}
        extraction_mode = ((profile.metadata or {}).get("extraction_mode") if profile is not None else None) or "standard"

        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows())
            values = [[_format_excel_cell(cell) for cell in row] for row in rows]
            rows_as_text = [" | ".join("" if cell is None else str(cell) for cell in row) for row in values[:20]]
            text_fragments.extend(rows_as_text)
            pages.append(PageExtraction(sheet_name=sheet.title, extracted_text="\n".join(rows_as_text)))

            header_map, header_row_index = _detect_header_map(values)
            if (header_row_index is None or header_map is None) and extraction_mode == "flexible":
                header_map, header_row_index = _infer_flexible_header_map(values)
            if header_row_index is None or header_map is None:
                continue

            sheet_items, row_stats = _extract_sheet_candidates(
                values,
                sheet_name=sheet.title,
                header_map=header_map,
                header_row_index=header_row_index,
                extraction_mode=extraction_mode,
            )
            sheet_stats[sheet.title] = row_stats
            items.extend(sheet_items)

        items = _filter_to_relevant_sheet_cluster(items)
        items, official_sheet = _annotate_sheet_roles(items)

        file_name_for_classification = profile.file_name if profile is not None else file_path.name
        document_kind = classify_document_from_text(file_name_for_classification, "\n".join(text_fragments))
        return ExtractionBatch(
            document_kind=document_kind,
            pages=pages,
            items=items,
            metadata={
                "sheet_count": len(workbook.worksheets),
                "extraction_mode": extraction_mode,
                "official_sheet": official_sheet,
                "sheet_stats": sheet_stats,
            },
        )
