from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook

from procurement_api.db import SessionLocal, init_db
from procurement_api.services import analyze_process, create_process, register_document
from procurement_core.schemas.process import ProcessCreate


ROOT = Path(__file__).resolve().parent
SAMPLE_DIR = ROOT / "sample_process"


def build_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Presupuesto"
    sheet.append(["Item", "Cantidad", "Unidad", "Valor Unitario", "Valor Total", "Lote"])
    sheet.append(["Guantes de nitrilo talla M caja x 100", 10, "caja", 32000, 320000, "Lote 1"])
    sheet.append(["Tapabocas quirúrgico caja x 50", 20, "caja", 16500, 330000, "Lote 1"])
    workbook.save(path)


def build_docx(path: Path) -> None:
    document = Document()
    document.add_heading("Estudio previo", level=1)
    document.add_paragraph("Proceso de adquisición de insumos médicos para atención primaria.")
    document.add_paragraph("Lote 1: Elementos de protección personal.")
    table = document.add_table(rows=1, cols=5)
    headers = ["Descripción", "Cantidad", "Unidad", "Valor unitario", "Valor total"]
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    row = table.add_row().cells
    row[0].text = "Guantes de nitrilo talla M caja x 100"
    row[1].text = "10"
    row[2].text = "caja"
    row[3].text = "32000"
    row[4].text = "320000"
    document.save(path)


def main() -> None:
    SAMPLE_DIR.mkdir(parents=True, exist_ok=True)
    xlsx_path = SAMPLE_DIR / "presupuesto_oficial_demo.xlsx"
    docx_path = SAMPLE_DIR / "estudio_previo_demo.docx"
    build_xlsx(xlsx_path)
    build_docx(docx_path)

    init_db()
    session = SessionLocal()
    try:
        process = create_process(
            session,
            ProcessCreate(
                external_reference="DEMO-001",
                name="Proceso demo de insumos médicos",
                contracting_entity="ESE Municipal Demo",
                description="Proceso generado automáticamente para pruebas end-to-end.",
                source_url="https://demo.entidad.gov.co/proceso/001",
            ),
        )
        register_document(session, process.id, xlsx_path.name, xlsx_path.read_bytes())
        register_document(session, process.id, docx_path.name, docx_path.read_bytes())
        session.commit()
        result = analyze_process(session, process.id)
        session.commit()
        print(f"Demo seeded successfully: {result}")
    finally:
        session.close()


if __name__ == "__main__":
    main()

